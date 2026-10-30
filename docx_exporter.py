"""
Word (.docx) MEB Sınav Kağıdı, Cevap Anahtarı ve Kitap Doğrulama Raporu Üretici
python-docx kullanarak A4 boyutunda, MEB ve Tevfik İleri İHO standartlarında yazılı kağıdı üretir.
"""

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from config import EXPORTS_DIR, get_config

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Hücre içi kenar boşluklarını ayarlar."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """Hücre kenarlıklarını belirler."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'CCCCCC'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def to_turkish_upper(text: str) -> str:
    """Türkçe karakter kurallarına uygun büyütme fonksiyonu (i -> İ, ı -> I)."""
    if not text:
        return ""
    s = str(text)
    # Varsa önceki bozuk İngilizce dönüşümleri onar
    s = s.replace("DIN KÜLTÜRÜ", "DİN KÜLTÜRÜ").replace("BILGISI", "BİLGİSİ").replace("AHIRET", "AHİRET")
    tr_map = str.maketrans({'i': 'İ', 'ı': 'I'})
    return s.translate(tr_map).upper()

def _calc_optimal_lines(q: dict) -> int:
    """Cevap anahtarı uzunluğuna göre dengeli satır sayısı hesaplar."""
    opts = q.get("options", [])
    if opts and isinstance(opts, list) and len(opts) > 0:
        return 0
    q_type = str(q.get("question_type") or "")
    q_txt = (q.get("question_text") or "")
    # BEP yapılandırılmış sorularda (Eşleştirme, D/Y, Kelime Havuzlu Boşluk Doldurma) ekstra cevap çizgisine gerek yoktur
    if q_type in ("bep_eslestirme", "bep_dogru_yanlis", "bep_bosluk_doldurma", "bep_coktan_secmeli_3") or "KELİME HAVUZU" in q_txt or "( ..... )" in q_txt or ".... ( a )" in q_txt:
        return 0
    ans = (q.get("answer_text") or "").strip()
    newline_count = ans.count("\n") + 1
    if len(ans) > 300 or newline_count >= 4:
        return 4
    return 3

def generate_exam_docx(exam_data: dict, layout: str = "2_column") -> Path:
    """MEB formatında öğrenci sınav kağıdı Word (.docx) dosyasını üretir."""
    doc = Document()
    cfg = get_config()
    
    # A4 Kenar Boşlukları (Dar kenar boşlukları - 1.27 cm / 0.5 inç)
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.page_width = Inches(8.27)  # A4
        section.page_height = Inches(11.69)
        
    header_info = exam_data.get("header_info", {})
    school_header = header_info.get("school_header") or cfg.get("school_header", "T.C. MİLLÎ EĞİTİM BAKANLIĞI")
    school_name = header_info.get("school_name") or cfg.get("school_name", "KOCAELİ / DARICA - Tevfik İleri İmam Hatip Ortaokulu")
    
    subject_title = exam_data.get("subject_name", exam_data.get("subject", "Din Kültürü ve Ahlak Bilgisi"))
    grade = exam_data.get("grade", 5)
    term = exam_data.get("term", "1. Dönem")
    exam_no = exam_data.get("exam_no", "1. Yazılı")
    scenario = exam_data.get("scenario", "Senaryo 1")
    exam_type = exam_data.get("exam_type", "yazili_sinav")
    total_points = exam_data.get("total_points", 100)
    show_scenario_in_header = exam_data.get("show_scenario_in_header", True)
    show_answer_lines = exam_data.get("show_answer_lines", True)
    bep_opts = exam_data.get("bep_options") or (header_info.get("bep_options") if isinstance(header_info, dict) else {}) or {}
    is_bep = (exam_type == "bep_sinav")

    # MEB BEP Yönetmeliği: BEP sınavları her zaman TEK SÜTUN olmak zorundadır
    if is_bep:
        layout = "1_column"

    if is_bep:
        show_bep_hdr = bep_opts.get("show_bep_header", True)
        if show_bep_hdr:
            sub_heading = f"{to_turkish_upper(term)} {to_turkish_upper(exam_no)} — BEP KAPSAMINDAKİ ÖĞRENCİ SINAVI"
        else:
            sub_heading = f"{to_turkish_upper(term)} {to_turkish_upper(exam_no)} SINAVI"
    elif exam_type == "etut_calisma":
        sub_heading = "ÇALIŞMA VE ETÜT KAĞIDI"
    elif exam_type == "yaprak_test":
        sub_heading = "KAZANIM DEĞERLENDİRME TESTİ"
    else:
        scenario_part = f" ({to_turkish_upper(scenario)})" if (show_scenario_in_header and scenario) else ""
        sub_heading = f"{to_turkish_upper(term)} {to_turkish_upper(exam_no)} SINAVI{scenario_part}"
    
    # 1. Başlık Anteti
    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_hdr.add_run(f"{to_turkish_upper(school_header)}\n")
    r1.bold = True
    r1.font.size = Pt(11 if is_bep else 10)
    r1.font.name = "Arial"
    
    r2 = p_hdr.add_run(f"{to_turkish_upper(school_name)}\n")
    r2.bold = True
    r2.font.size = Pt(11.5 if is_bep else 10.5)
    r2.font.name = "Arial"
    
    r3 = p_hdr.add_run(f"2026-2027 EĞİTİM ÖĞRETİM YILI {grade}. SINIF {to_turkish_upper(subject_title)} DERSİ\n")
    r3.bold = True
    r3.font.size = Pt(11 if is_bep else 10)
    r3.font.name = "Arial"
    
    r4 = p_hdr.add_run(sub_heading)
    r4.bold = True
    r4.font.size = Pt(12 if is_bep else 10.5)
    if is_bep:
        r4.font.color.rgb = RGBColor(30, 58, 138)
    r4.font.name = "Arial"
    p_hdr.paragraph_format.space_after = Pt(6 if is_bep else 4)
    
    # 2. Öğrenci Bilgi Tablosu
    info_table = doc.add_table(rows=2, cols=4)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    
    col_widths = [Inches(1.2), Inches(3.2), Inches(1.5), Inches(1.3)]
    for row in info_table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width
            set_cell_margins(row.cells[idx], top=60, bottom=60, left=100, right=100)
            set_cell_border(row.cells[idx], 
                            top={'val': 'single', 'sz': 4, 'color': '888888'},
                            bottom={'val': 'single', 'sz': 4, 'color': '888888'},
                            left={'val': 'single', 'sz': 4, 'color': '888888'},
                            right={'val': 'single', 'sz': 4, 'color': '888888'})
                            
    # 1. Satır: Adı Soyadı & Puan Kutusu
    c0 = info_table.rows[0].cells[0]
    p = c0.paragraphs[0]
    r = p.add_run("Adı ve Soyadı:")
    r.bold = True
    r.font.size = Pt(9)
    r.font.name = "Arial"
    
    c1 = info_table.rows[0].cells[1]
    p = c1.paragraphs[0]
    p.add_run(".............................................................")
    
    c2 = info_table.rows[0].cells[2]
    p = c2.paragraphs[0]
    r = p.add_run("ALDIĞI PUAN:")
    r.bold = True
    r.font.size = Pt(9)
    r.font.name = "Arial"
    
    c3 = info_table.rows[0].cells[3]
    p = c3.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"....... / {total_points}")
    r.bold = True
    r.font.size = Pt(10)
    r.font.name = "Arial"
    
    # 2. Satır: Sınıf/Şube & No & Tarih
    c0 = info_table.rows[1].cells[0]
    p = c0.paragraphs[0]
    r = p.add_run("Sınıfı / Şubesi:")
    r.bold = True
    r.font.size = Pt(9)
    r.font.name = "Arial"
    
    c1 = info_table.rows[1].cells[1]
    p = c1.paragraphs[0]
    p.add_run(f"{grade} / .....       Okul No: ...............")
    
    c2 = info_table.rows[1].cells[2]
    p = c2.paragraphs[0]
    r = p.add_run("Tarih / İmza:")
    r.bold = True
    r.font.size = Pt(9)
    r.font.name = "Arial"
    
    c3 = info_table.rows[1].cells[3]
    p = c3.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("...../...../2026")
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    questions = exam_data.get("questions", [])
    
    # Düzen: 2 Sütunlu (2_column) veya Tek Sütunlu (1_column)
    if layout == "2_column" and len(questions) > 1:
        mid = (len(questions) + 1) // 2
        col1_questions = questions[:mid]
        col2_questions = questions[mid:]
        
        main_table = doc.add_table(rows=1, cols=2)
        main_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        main_table.autofit = False
        
        c_left = main_table.rows[0].cells[0]
        c_right = main_table.rows[0].cells[1]
        
        c_left.width = Inches(3.6)
        c_right.width = Inches(3.6)
        
        set_cell_border(c_left, right={'val': 'dashed', 'sz': 6, 'color': 'AAAAAA'})
        set_cell_margins(c_left, top=50, bottom=50, left=50, right=100)
        set_cell_margins(c_right, top=50, bottom=50, left=100, right=50)
        
        def render_questions_into_cell(cell, q_list):
            for q in q_list:
                p_q = cell.add_paragraph()
                r_num = p_q.add_run(f"Soru {q.get('q_order')}) ")
                r_num.bold = True
                r_num.font.size = Pt(9.5)
                r_num.font.name = "Arial"
                
                r_pts = p_q.add_run(f"({q.get('points', 10)} Puan) ")
                r_pts.bold = True
                r_pts.font.size = Pt(8.5)
                r_pts.font.color.rgb = RGBColor(100, 100, 100)
                r_pts.font.name = "Arial"
                
                p_q.paragraph_format.space_after = Pt(2)
                p_q.paragraph_format.line_spacing = 1.15
                
                ar_txt = (q.get("arabic_text") or "").strip()
                if ar_txt:
                    p_ar = cell.add_paragraph()
                    p_ar.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r_ar = p_ar.add_run(ar_txt)
                    r_ar.font.name = "Traditional Arabic"
                    r_ar.font.size = Pt(14)
                    r_ar.bold = True
                    p_ar.paragraph_format.space_after = Pt(2)
                    
                p_text = cell.add_paragraph()
                r_txt = p_text.add_run(q.get("question_text", ""))
                r_txt.font.size = Pt(9)
                r_txt.font.name = "Arial"
                p_text.paragraph_format.space_after = Pt(4)
                
                opts = q.get("options", [])
                if opts and isinstance(opts, list) and len(opts) > 0:
                    p_opts = cell.add_paragraph()
                    p_opts.paragraph_format.space_after = Pt(10)
                    p_opts.paragraph_format.line_spacing = 1.25
                    r_op = p_opts.add_run("\n".join(str(o) for o in opts))
                    r_op.font.size = Pt(8.5)
                    r_op.font.name = "Arial"
                else:
                    if show_answer_lines:
                        num_lines = _calc_optimal_lines(q)
                        p_lines = cell.add_paragraph()
                        p_lines.paragraph_format.line_spacing = 1.3
                        p_lines.paragraph_format.space_after = Pt(10)
                        lines_str = "Cevap:\n" + ("................................................................................\n" * num_lines)
                        r_ln = p_lines.add_run(lines_str.rstrip("\n"))
                        r_ln.font.size = Pt(8)
                        r_ln.font.color.rgb = RGBColor(160, 160, 160)
                    else:
                        p_lines = cell.add_paragraph()
                        p_lines.paragraph_format.space_after = Pt(36)
                        r_ln = p_lines.add_run("Cevap:")
                        r_ln.font.size = Pt(8)
                        r_ln.font.color.rgb = RGBColor(140, 140, 140)

        render_questions_into_cell(c_left, col1_questions)
        render_questions_into_cell(c_right, col2_questions)
        
    else:
        # Tek Sütunlu Sayfa Düzeni (veya BEP Özel Eğitim Tek Sütun Düzeni)
        bep_font_choice = bep_opts.get("font_size", "13pt")
        bep_q_style = bep_opts.get("question_style", "karma")
        is_large_bep = (is_bep and bep_font_choice == "16pt")
        base_q_pt = 14 if is_large_bep else (11.5 if is_bep else 9.5)
        badge_pt = 14 if is_large_bep else (12 if is_bep else 10)
        ar_pt = 20 if is_large_bep else (16.5 if is_bep else 15)
        line_spc = 1.4 if is_bep else 1.15

        for q in questions:
            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(4 if is_bep else 2)
            p_q.paragraph_format.space_after = Pt(1.5)
            r_num = p_q.add_run(f"Soru {q.get('q_order')}) ")
            r_num.bold = True
            r_num.font.size = Pt(badge_pt)
            r_num.font.name = "Arial"
            
            r_pts = p_q.add_run(f"({q.get('points', 10)} Puan)")
            r_pts.bold = True
            r_pts.font.size = Pt(badge_pt - 1)
            r_pts.font.color.rgb = RGBColor(30, 58, 138) if is_bep else RGBColor(100, 100, 100)
            r_pts.font.name = "Arial"
            
            ar_txt = (q.get("arabic_text") or "").strip()
            if ar_txt:
                p_ar = doc.add_paragraph()
                p_ar.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_ar.paragraph_format.space_after = Pt(2)
                r_ar = p_ar.add_run(ar_txt)
                r_ar.font.name = "Traditional Arabic"
                r_ar.font.size = Pt(ar_pt)
                r_ar.bold = True
                
            raw_q_txt = str(q.get("question_text", ""))
            if is_bep:
                while "\n\n" in raw_q_txt:
                    raw_q_txt = raw_q_txt.replace("\n\n", "\n")
                raw_q_txt = raw_q_txt.strip()

            p_text = doc.add_paragraph()
            p_text.paragraph_format.line_spacing = line_spc
            p_text.paragraph_format.space_after = Pt(3 if is_bep else 4)
            r_txt = p_text.add_run(raw_q_txt)
            r_txt.font.size = Pt(base_q_pt)
            r_txt.font.name = "Arial"
            
            opts = q.get("options", [])
            q_type_str = str(q.get("question_type") or "")
            if is_bep and bep_q_style != "coktan_secmeli_3" and (
                q_type_str in ("bep_eslestirme", "bep_dogru_yanlis", "bep_bosluk_doldurma")
                or "KELİME HAVUZU" in raw_q_txt
                or "( ..... )" in raw_q_txt
                or ".... ( a )" in raw_q_txt
            ):
                opts = []

            if opts and isinstance(opts, list) and len(opts) > 0:
                p_opts = doc.add_paragraph()
                p_opts.paragraph_format.line_spacing = 1.35 if is_bep else 1.25
                p_opts.paragraph_format.space_after = Pt(8 if is_bep else 12)
                short_3col = (len(opts) == 3 and all(len(str(o).strip()) <= 36 for o in opts))
                opts_str = ("        ".join(str(o).strip() for o in opts)) if short_3col else ("\n".join(str(o) for o in opts))
                r_op = p_opts.add_run(opts_str)
                r_op.font.size = Pt(base_q_pt)
                r_op.font.bold = True if is_bep else False
                r_op.font.name = "Arial"
            else:
                num_lines = _calc_optimal_lines(q)
                if num_lines > 0:
                    if show_answer_lines:
                        p_lines = doc.add_paragraph()
                        p_lines.paragraph_format.line_spacing = 1.4 if is_bep else 1.25
                        p_lines.paragraph_format.space_after = Pt(10)
                        lines_str = "Cevap:\n" + ("...............................................................................................................................\n" * num_lines)
                        r_ln = p_lines.add_run(lines_str.rstrip("\n"))
                        r_ln.font.size = Pt(10.5 if is_bep else 8.5)
                        r_ln.font.color.rgb = RGBColor(160, 160, 160)
                    else:
                        p_lines = doc.add_paragraph()
                        p_lines.paragraph_format.space_after = Pt(36)
                        r_ln = p_lines.add_run("Cevap:")
                        r_ln.font.size = Pt(10.5 if is_bep else 8.5)
                        r_ln.font.color.rgb = RGBColor(140, 140, 140)
                elif is_bep:
                    p_sep = doc.add_paragraph()
                    p_sep.paragraph_format.space_after = Pt(4)
            
    # Alt Not
    p_ft = doc.add_paragraph()
    p_ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if is_bep:
        accom = bep_opts.get("accommodation", "ek_sure")
        accom_str = "Okuyucu/Kodlayıcı Desteği & +%25-%50 Ek Süre" if accom == "okuyucu_kodlayici" else "+%25-%50 BEP Ek Süresi"
        ft_text = f"BEP Kapsamındaki Öğrenci Sınavı ({accom_str}) • Yazım ve noktalama hatalarından puan kırılmaz. Başarılar dilerim."
    else:
        ft_text = "Sınav süresi 40 dakikadır. Başarılar dilerim."
    r_ft = p_ft.add_run(ft_text)
    r_ft.font.italic = True
    r_ft.font.size = Pt(9.5 if is_bep else 8.5)
    r_ft.font.name = "Arial"
    
    file_id = exam_data.get("id", "yeni")
    bep_prefix = "BEP_" if is_bep else ""
    filename = f"{bep_prefix}Sinav_Kagidi_{subject_title}_{grade}Sinif_{term}_{file_id}.docx".replace(" ", "_")
    output_path = EXPORTS_DIR / filename
    doc.save(str(output_path))
    return output_path

def generate_answer_key_docx(exam_data: dict) -> Path:
    """Öğretmen Cevap Anahtarı ve Puanlama Rubriği Word dosyasını üretir."""
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        
    subject_title = exam_data.get("subject_name", exam_data.get("subject", "Din Kültürü ve Ahlak Bilgisi"))
    grade = exam_data.get("grade", 5)
    term = exam_data.get("term", "1. Dönem")
    exam_no = exam_data.get("exam_no", "1. Yazılı")
    scenario = exam_data.get("scenario", "Senaryo 1")
    exam_type = exam_data.get("exam_type", "yazili_sinav")
    is_bep = (exam_type == "bep_sinav")
    
    cfg = get_config()
    school_title = exam_data.get("header_info", {}).get("school_name") or cfg.get("school_name", "KOCAELİ / DARICA - Tevfik İleri İmam Hatip Ortaokulu")
    
    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_hdr.add_run(f"{to_turkish_upper(school_title)}\n")
    r1.bold = True
    r1.font.size = Pt(11)
    
    show_scenario_in_header = exam_data.get("show_scenario_in_header", True)
    if is_bep:
        r2 = p_hdr.add_run(f"{grade}. SINIF {to_turkish_upper(subject_title)} {to_turkish_upper(term)} {to_turkish_upper(exam_no)}\nBEP KAPSAMINDAKİ ÖĞRENCİ SINAVI — CEVAP ANAHTARI VE DEĞERLENDİRME RUBRİĞİ")
    else:
        scenario_part = f" ({to_turkish_upper(scenario)})" if (show_scenario_in_header and scenario) else ""
        r2 = p_hdr.add_run(f"{grade}. SINIF {to_turkish_upper(subject_title)} {to_turkish_upper(term)} {to_turkish_upper(exam_no)}{scenario_part}\nCEVAP ANAHTARI VE DERECELİ PUANLAMA ANAHTARI (RUBRİK)")
    r2.bold = True
    r2.font.size = Pt(11.5)
    r2.font.color.rgb = RGBColor(0, 51, 102)
    p_hdr.paragraph_format.space_after = Pt(8)

    if is_bep:
        p_bep_note = doc.add_paragraph()
        r_bn = p_bep_note.add_run(
            "MEB Özel Eğitim Hizmetleri Yönetmeliği BEP Puanlama ve Evrak Esasları:\n"
            "• Yazım hataları, noktalama işaretleri veya bozuk el yazısı nedeniyle puan kırılmaz.\n"
            "• Öğrencinin sorunun özünü kavradığını gösteren her türlü doğru ifadeye ve işaretlemeye tam puan verilir.\n"
            "• Öğrencinin sınav kağıdı zümre BEP dosyasında saklanır ve e-Okul sistemine BEP statüsünde 100 üzerinden girilir."
        )
        r_bn.font.size = Pt(9.5)
        r_bn.font.color.rgb = RGBColor(30, 58, 138)
        p_bep_note.paragraph_format.space_after = Pt(10)
    
    questions = exam_data.get("questions", [])
    for q in questions:
        p_q = doc.add_paragraph()
        r_num = p_q.add_run(f"SORU {q.get('q_order')} [Toplam: {q.get('points', 10)} Puan] ")
        r_num.bold = True
        r_num.font.size = Pt(10)
        
        outcome = q.get("learning_outcome", "")
        if outcome:
            r_out = p_q.add_run(f" | Kazanım: {outcome}")
            r_out.font.italic = True
            r_out.font.size = Pt(8.5)
            r_out.font.color.rgb = RGBColor(120, 120, 120)
            
        p_q.paragraph_format.space_after = Pt(2)
        
        p_txt = doc.add_paragraph()
        r_txt = p_txt.add_run(q.get("question_text", ""))
        r_txt.font.size = Pt(9.5)
        p_txt.paragraph_format.space_after = Pt(4)
        
        p_ans = doc.add_paragraph()
        r_at = p_ans.add_run("Beklenen Örnek Cevap:\n")
        r_at.bold = True
        r_at.font.size = Pt(9.5)
        r_at.font.color.rgb = RGBColor(0, 102, 51)
        r_ac = p_ans.add_run(q.get("answer_text", ""))
        r_ac.font.size = Pt(9.5)
        p_ans.paragraph_format.space_after = Pt(4)
        
        rubric = q.get("rubric", [])
        if rubric and isinstance(rubric, list):
            rubric_table = doc.add_table(rows=1, cols=2)
            rubric_table.autofit = False
            rubric_table.rows[0].cells[0].width = Inches(6.0)
            rubric_table.rows[0].cells[1].width = Inches(1.2)
            
            h0 = rubric_table.rows[0].cells[0]
            h0.paragraphs[0].add_run("Puanlama Kriteri (Kısmi Puan Basamakları)").bold = True
            h1 = rubric_table.rows[0].cells[1]
            h1.paragraphs[0].add_run("Puan").bold = True
            h1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            set_cell_border(h0, bottom={'val': 'single', 'sz': 4, 'color': '888888'})
            set_cell_border(h1, bottom={'val': 'single', 'sz': 4, 'color': '888888'})
            
            for item in rubric:
                row = rubric_table.add_row()
                row.cells[0].width = Inches(6.0)
                row.cells[1].width = Inches(1.2)
                
                c_text = item.get("criterion", "") if isinstance(item, dict) else str(item)
                c_pts = item.get("points", "") if isinstance(item, dict) else ""
                
                row.cells[0].paragraphs[0].add_run(f"• {c_text}").font.size = Pt(9)
                p_p = row.cells[1].paragraphs[0]
                p_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_p.add_run(f"{c_pts} Puan" if c_pts else "").font.size = Pt(9)
                
                set_cell_margins(row.cells[0], top=40, bottom=40, left=60, right=60)
                set_cell_margins(row.cells[1], top=40, bottom=40, left=60, right=60)
                set_cell_border(row.cells[0], bottom={'val': 'single', 'sz': 2, 'color': 'DDDDDD'})
                set_cell_border(row.cells[1], bottom={'val': 'single', 'sz': 2, 'color': 'DDDDDD'})
                
        doc.add_paragraph().paragraph_format.space_after = Pt(8)
        
    file_id = exam_data.get("id", "yeni")
    filename = f"Cevap_Anahtari_{subject_title}_{grade}Sinif_{term}_{file_id}.docx".replace(" ", "_")
    output_path = EXPORTS_DIR / filename
    doc.save(str(output_path))
    return output_path

def generate_source_verification_docx(exam_data: dict) -> Path:
    """Soruların MEB Ders Kitabı İlişkisi ve Pedagojik Doğrulama Raporu Word dosyasını üretir."""
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    subject_title = exam_data.get("subject_name", exam_data.get("subject", "Din Kültürü ve Ahlak Bilgisi"))
    grade = exam_data.get("grade", 5)
    term = exam_data.get("term", "1. Dönem")
    exam_no = exam_data.get("exam_no", "1. Yazılı")

    cfg = get_config()
    school_title = exam_data.get("header_info", {}).get("school_name") or cfg.get("school_name", "KOCAELİ / DARICA - Tevfik İleri İmam Hatip Ortaokulu")

    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_hdr.add_run(f"{to_turkish_upper(school_title)}\n")
    r1.bold = True
    r1.font.size = Pt(11)

    r2 = p_hdr.add_run(f"{grade}. SINIF {to_turkish_upper(subject_title)} {to_turkish_upper(term)} {to_turkish_upper(exam_no)}\nKİTAP İLİŞKİSİ, SAYFA EŞLEŞTİRME VE PEDAGOJİK DOĞRULAMA RAPORU")
    r2.bold = True
    r2.font.size = Pt(11)
    r2.font.color.rgb = RGBColor(146, 64, 14)
    p_hdr.paragraph_format.space_after = Pt(10)

    questions = exam_data.get("questions", [])
    for q in questions:
        p_q = doc.add_paragraph()
        r_num = p_q.add_run(
            f"SORU {q.get('q_order')} ({q.get('points', 10)} Puan) | "
            f"Bilişsel Düzey: {q.get('cognitive_level', 'Analiz')} | "
            f"Güçlük: {q.get('difficulty', 'Orta')}\n"
        )
        r_num.bold = True
        r_num.font.size = Pt(10)
        r_num.font.color.rgb = RGBColor(30, 58, 138)

        r_src = p_q.add_run(
            f"Kaynak Kitap: {q.get('source_book', f'{grade}. Sınıf MEB Ders Kitabı')} — "
            f"Sayfa / Bölüm: {q.get('source_page', 'İlgili Ünite')}\n"
            f"MEB Kazanımı: {q.get('learning_outcome', 'Genel Kazanım')}"
        )
        r_src.bold = True
        r_src.font.size = Pt(9)
        r_src.font.color.rgb = RGBColor(71, 85, 105)
        p_q.paragraph_format.space_after = Pt(3)

        p_txt = doc.add_paragraph()
        p_txt.add_run("Soru Metni: ").bold = True
        r_t = p_txt.add_run(q.get("question_text", ""))
        r_t.font.size = Pt(9.5)
        p_txt.paragraph_format.space_after = Pt(3)

        p_ctx = doc.add_paragraph()
        r_cl = p_ctx.add_run("Kitaptaki Pedagojik Bağlam ve Sayfa Özeti: ")
        r_cl.bold = True
        r_cl.font.size = Pt(9)
        r_cl.font.color.rgb = RGBColor(146, 64, 14)
        r_cv = p_ctx.add_run(q.get("source_context", "MEB ders kitabı ünite temaları esas alınmıştır."))
        r_cv.font.size = Pt(9)
        p_ctx.paragraph_format.space_after = Pt(3)

        p_ans = doc.add_paragraph()
        r_al = p_ans.add_run("Beklenen Çözüm / Cevap: ")
        r_al.bold = True
        r_al.font.size = Pt(9)
        r_al.font.color.rgb = RGBColor(6, 95, 70)
        r_av = p_ans.add_run(q.get("answer_text", ""))
        r_av.font.size = Pt(9)
        p_ans.paragraph_format.space_after = Pt(10)

    file_id = exam_data.get("id", "yeni")
    filename = f"Kitap_Dogrulama_{subject_title}_{grade}Sinif_{term}_{file_id}.docx".replace(" ", "_")
    output_path = EXPORTS_DIR / filename
    doc.save(str(output_path))
    return output_path
