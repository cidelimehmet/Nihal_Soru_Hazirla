"""
FastAPI Ana Sunucu ve Genişletilmiş API Endpointleri
Tevfik İleri İHO Soru Hazırlama, Soru Bankası & Çoklu Çıktı Sistemi
"""

import os
from pathlib import Path
from fastapi import FastAPI, HTTPException, Request, Depends
from fastapi.responses import HTMLResponse, FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import uvicorn

from config import get_config, save_config, BASE_DIR, EXPORTS_DIR, DATA_DIR
from curriculum_data import get_subjects, get_grades, get_units, CURRICULUM
from gemini_engine import generate_questions_with_gemini
from docx_exporter import generate_exam_docx, generate_answer_key_docx
import database

app = FastAPI(title="Tevfik İleri İHO Soru Hazırlama & Soru Bankası Sistemi", version="2.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

STATIC_DIR = BASE_DIR / "static"
STATIC_DIR.mkdir(exist_ok=True)
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")

# Pydantic Modelleri
class LoginRequest(BaseModel):
    username: str
    password: str

class GenerateExamRequest(BaseModel):
    subject: str
    grade: int
    exam_type: str = "yazili_sinav" # 'yazili_sinav', 'etut_calisma', 'yaprak_test'
    term: str = "1. Dönem"
    exam_no: str = "1. Yazılı"
    scenario: str = "Senaryo 1"
    question_count: int = 6
    selected_units: list[str] = []
    difficulty: str = "Dengeli"

class SingleQuestionRegenRequest(BaseModel):
    subject: str
    grade: int
    term: str = "1. Dönem"
    exam_no: str = "1. Yazılı"
    q_order: int
    selected_units: list[str] = []
    exam_type: str = "yazili_sinav"

class SaveExamRequest(BaseModel):
    id: int | None = None
    title: str = "Yeni Sınav"
    exam_type: str = "yazili_sinav"
    subject: str
    grade: int
    term: str
    exam_no: str
    scenario: str
    layout: str = "2_column"
    total_points: int = 100
    questions: list
    header_info: dict = {}

class BankGenerateRequest(BaseModel):
    subject: str
    grade: int
    count: int = 5
    selected_units: list[str] = []
    difficulty: str = "Dengeli"
    question_type: str = "acik_uclu"

class ConfigUpdateRequest(BaseModel):
    gemini_api_key: str | None = None
    gemini_model: str | None = None
    school_header: str | None = None
    school_name: str | None = None
    teacher_name: str | None = None
    login_username: str | None = None
    login_password: str | None = None

# --- ENDPOINTS ---

@app.get("/", response_class=HTMLResponse)
async def serve_index():
    index_file = STATIC_DIR / "index.html"
    if index_file.exists():
        with open(index_file, "r", encoding="utf-8") as f:
            return HTMLResponse(content=f.read())
    return HTMLResponse(content="<h1>Uygulama yükleniyor...</h1>")

@app.get("/api/config")
async def api_get_config():
    cfg = get_config()
    safe_cfg = dict(cfg)
    if safe_cfg.get("gemini_api_key"):
        k = safe_cfg["gemini_api_key"]
        safe_cfg["has_api_key"] = True
        safe_cfg["gemini_api_key_masked"] = k[:6] + "..." + k[-4:] if len(k) > 10 else "***"
    else:
        safe_cfg["has_api_key"] = False
        safe_cfg["gemini_api_key_masked"] = ""
    safe_cfg.pop("login_password", None)
    return safe_cfg

@app.post("/api/config")
async def api_save_config(data: ConfigUpdateRequest):
    update_data = {k: v for k, v in data.dict().items() if v is not None}
    save_config(update_data)
    return {"status": "success", "message": "Ayarlar kaydedildi."}

@app.post("/api/login")
async def api_login(req: LoginRequest):
    cfg = get_config()
    if req.username == cfg["login_username"] and req.password == cfg["login_password"]:
        return {
            "status": "success",
            "user": {
                "username": req.username,
                "teacher_name": cfg["teacher_name"],
                "school_name": cfg["school_name"]
            }
        }
    raise HTTPException(status_code=401, detail="Kullanıcı adı veya şifre hatalı!")

@app.get("/api/curriculum")
async def api_get_curriculum():
    from curriculum_data import load_curriculum_dict
    return load_curriculum_dict()

@app.get("/api/units")
async def api_get_units(subject: str, grade: int):
    units = get_units(subject, grade)
    return {"units": units}

# --- MÜFREDAT VE KONU YÖNETİCİSİ ENDPOINTLERİ (YENİ TALEP) ---
@app.get("/api/curriculum-manager")
async def api_get_curriculum_manager(subject: str, grade: int):
    from curriculum_data import get_units, load_curriculum_dict
    curr = load_curriculum_dict()
    units = get_units(subject, grade)
    return {
        "status": "success",
        "subject": subject,
        "grade": int(grade),
        "subject_name": curr.get(subject, {}).get("name", subject),
        "units": units
    }

@app.post("/api/curriculum-manager/save")
async def api_save_curriculum_manager(req: dict):
    from curriculum_data import load_curriculum_dict, save_curriculum_dict
    subject = req.get("subject")
    grade = int(req.get("grade", 5))
    new_units = req.get("units", [])
    
    curr = load_curriculum_dict()
    if subject in curr and grade in curr[subject]["grades"]:
        curr[subject]["grades"][grade]["units"] = new_units
        save_curriculum_dict(curr)
        return {"status": "success", "message": "Üniteler ve konular başarıyla güncellendi!"}
    raise HTTPException(status_code=400, detail="Geçersiz ders veya sınıf.")

@app.post("/api/curriculum-manager/reset")
async def api_reset_curriculum_manager():
    from curriculum_data import reset_curriculum_to_default
    reset_curriculum_to_default()
    return {"status": "success", "message": "Müfredat MEB orijinal varsayılan ayarlarına döndürüldü."}

@app.post("/api/generate-exam")
async def api_generate_exam(req: GenerateExamRequest):
    try:
        data = generate_questions_with_gemini(
            subject_id=req.subject,
            grade=req.grade,
            term=req.term,
            exam_no=req.exam_no,
            scenario=req.scenario,
            question_count=req.question_count,
            selected_unit_ids=req.selected_units,
            difficulty=req.difficulty,
            exam_type=req.exam_type
        )
        return {"status": "success", "data": data}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Soru hazırlama hatası: {str(e)}")

@app.post("/api/regenerate-question")
async def api_regenerate_question(req: SingleQuestionRegenRequest):
    """Tek bir soruyu yeniden üretir."""
    try:
        data = generate_questions_with_gemini(
            subject_id=req.subject,
            grade=req.grade,
            term=req.term,
            exam_no=req.exam_no,
            scenario="Özel",
            question_count=1,
            selected_unit_ids=req.selected_units,
            difficulty="Dengeli",
            exam_type=req.exam_type
        )
        if data.get("questions"):
            q = data["questions"][0]
            q["q_order"] = req.q_order
            return {"status": "success", "question": q}
        raise HTTPException(status_code=500, detail="Soru üretilemedi.")
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/save-exam")
async def api_save_exam(req: SaveExamRequest):
    cfg = get_config()
    exam_dict = req.dict()
    if not exam_dict.get("header_info"):
        exam_dict["header_info"] = {
            "school_header": cfg.get("school_header"),
            "school_name": cfg.get("school_name"),
            "teacher_name": cfg.get("teacher_name")
        }
    exam_id = database.save_exam(exam_dict)
    return {"status": "success", "exam_id": exam_id, "message": "Sınav başarıyla kaydedildi."}

@app.get("/api/exams")
async def api_list_exams():
    exams = database.list_exams()
    return {"exams": exams}

@app.get("/api/exams/{exam_id}")
async def api_get_exam(exam_id: int):
    exam = database.get_exam(exam_id)
    if not exam:
        raise HTTPException(status_code=404, detail="Sınav bulunamadı.")
    return {"exam": exam}

@app.delete("/api/exams/{exam_id}")
async def api_delete_exam(exam_id: int):
    database.delete_exam(exam_id)
    return {"status": "success", "message": "Sınav silindi."}

# --- SORU BANKASI ENDPOINTLERİ ---

@app.get("/api/question-bank")
async def api_get_question_bank(
    subject: str | None = None,
    grade: int | None = None,
    difficulty: str | None = None,
    q_type: str | None = None,
    search: str | None = None
):
    items = database.get_question_bank_items(subject, grade, difficulty, q_type, search)
    return {"status": "success", "items": items}

@app.post("/api/question-bank")
async def api_add_question_to_bank(req: dict):
    new_id = database.add_question_to_bank(req)
    return {"status": "success", "id": new_id, "message": "Soru bankasına eklendi."}

@app.delete("/api/question-bank/{qid}")
async def api_delete_question_from_bank(qid: int):
    database.delete_question_from_bank(qid)
    return {"status": "success", "message": "Soru silindi."}

@app.post("/api/generate-bank-questions")
async def api_generate_bank_questions(req: BankGenerateRequest):
    """Sınav yapmadan kitaptan doğrudan soru bankasına soru üretip depolar."""
    try:
        data = generate_questions_with_gemini(
            subject_id=req.subject,
            grade=req.grade,
            term="1. Dönem",
            exam_no="Soru Havuzu",
            scenario="Özel",
            question_count=req.count,
            selected_unit_ids=req.selected_units,
            difficulty=req.difficulty,
            exam_type="yaprak_test" if req.question_type == "coktan_secmeli" else "yazili_sinav"
        )
        saved_count = 0
        for q in data.get("questions", []):
            q["subject"] = req.subject
            q["grade"] = req.grade
            database.add_question_to_bank(q)
            saved_count += 1
            
        return {"status": "success", "saved_count": saved_count, "message": f"{saved_count} soru soru bankasına eklendi."}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

# --- TASLAK VE OTOMATİK KAYIT ENDPOINTLERİ (15 SN VERİ KAYBI ÖNLEME) ---
DRAFT_FILE = DATA_DIR / "active_draft.json"

@app.get("/api/draft")
async def api_get_draft():
    if DRAFT_FILE.exists():
        try:
            import json
            with open(DRAFT_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                return {"status": "success", "draft": data}
        except Exception:
            pass
    return {"status": "none", "draft": None}

@app.post("/api/draft")
async def api_save_draft(req: dict):
    try:
        import json
        with open(DRAFT_FILE, "w", encoding="utf-8") as f:
            json.dump(req, f, ensure_ascii=False, indent=2)
        return {"status": "success", "message": "Taslak kaydedildi."}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.delete("/api/draft")
async def api_delete_draft():
    if DRAFT_FILE.exists():
        try:
            DRAFT_FILE.unlink()
        except Exception:
            pass
    return {"status": "success", "message": "Taslak temizlendi."}

# --- DIŞA AKTARMA (DOCX) ---

@app.post("/api/export/docx")
async def api_export_docx(req: dict):
    exam_data = req.get("exam_data", {})
    doc_type = req.get("type", "exam") # "exam", "answer_key", "source_verification"
    layout = req.get("layout", "2_column")
    
    subject_id = exam_data.get("subject", "dkab")
    subject_name = CURRICULUM.get(subject_id, {}).get("name", exam_data.get("subject", "Ders"))
    exam_data["subject_name"] = subject_name
    
    cfg = get_config()
    if "header_info" not in exam_data or not exam_data["header_info"]:
        exam_data["header_info"] = {
            "school_header": cfg.get("school_header"),
            "school_name": cfg.get("school_name")
        }
        
    if doc_type == "answer_key":
        file_path = generate_answer_key_docx(exam_data)
        download_name = f"Cevap_Anahtari_{subject_name}_{exam_data.get('grade')}Sinif.docx"
    elif doc_type == "source_verification":
        # Kitap ilişkisi doğrulama belgesi
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        doc = Document()
        for s in doc.sections:
            s.top_margin = Inches(0.5); s.bottom_margin = Inches(0.5)
            s.left_margin = Inches(0.5); s.right_margin = Inches(0.5)
        p = doc.add_paragraph()
        r = p.add_run(f"KİTAP İLİŞKİSİ VE PEDAGOJİK SORU DOĞRULAMA RAPORU\n{exam_data.get('grade')}. Sınıf {subject_name}\n")
        r.bold = True; r.font.size = Pt(12)
        for q in exam_data.get("questions", []):
            p_q = doc.add_paragraph()
            p_q.add_run(f"Soru {q.get('q_order')}: ").bold = True
            p_q.add_run(f"{q.get('question_text')}\n")
            p_q.add_run(f"• Kaynak Kitap: {q.get('source_book', 'Ders Kitabı')} | {q.get('source_page', '')}\n").italic = True
            p_q.add_run(f"• Bilişsel Düzey: {q.get('cognitive_level', 'Analiz')} | Güçlük: {q.get('difficulty', 'Orta')}\n")
            p_q.add_run(f"• Pedagojik Bağlam: {q.get('source_context', '')}\n")
            p_q.add_run(f"• Doğru Cevap: {q.get('answer_text', '')}\n").bold = True
        file_path = EXPORTS_DIR / f"Kitap_Dogrulama_{subject_name}_{exam_data.get('grade')}Sinif.docx"
        doc.save(str(file_path))
        download_name = f"Kitap_Dogrulama_{subject_name}_{exam_data.get('grade')}Sinif.docx"
    else:
        file_path = generate_exam_docx(exam_data, layout=layout)
        download_name = f"Sinav_Kagidi_{subject_name}_{exam_data.get('grade')}Sinif_{layout}.docx"
        
    return FileResponse(
        path=str(file_path),
        filename=download_name,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

if __name__ == "__main__":
    import threading
    import time
    import urllib.request
    import webbrowser

    cfg = get_config()
    port = cfg.get("port", 8000)

    # Sunucu tam ayağa kalktığında tarayıcıyı açan akıllı kontrol mekanizması
    def open_browser_when_ready():
        url = f"http://127.0.0.1:{port}"
        check_url = f"http://127.0.0.1:{port}/api/config"
        for _ in range(40): # En fazla 20 saniye kontrol et
            try:
                time.sleep(0.5)
                req = urllib.request.Request(check_url, headers={'User-Agent': 'Mozilla/5.0'})
                with urllib.request.urlopen(req, timeout=1) as resp:
                    if resp.status == 200:
                        print(f"Sunucu hazir! Tarayici aciliyor: {url}")
                        webbrowser.open(url)
                        break
            except Exception:
                pass

    threading.Thread(target=open_browser_when_ready, daemon=True).start()

    print(f"Tevfik İleri İHO Soru Hazırlama & Soru Bankası Başlatılıyor...")
    print(f"Adres: http://localhost:{port}")
    uvicorn.run("main:app", host="127.0.0.1", port=port, reload=True)
